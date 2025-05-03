import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


def save_json(data, filename):
    with open(filename, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4)


def load_json(filename):
    try:
        with open(filename, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return []


def write_to_csv(json_file, csv_file):
    try:
        with open(json_file, encoding="utf-8") as input_file:
            df = pd.read_json(input_file)

        df = df.rename(columns={
            "type": "Type",
            "name": "Name",
            "symbol": "Symbol",
            "shares": "Shares",
            "purchase_price": "Purchase Price",
            "sector": "Sector"
        })

        df.to_csv(csv_file, encoding="utf-8", index=False)
        print(f"Investments have been written to {csv_file} successfully.")

        excel_file = csv_file.replace(".csv", ".xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Investments"

        for row in dataframe_to_rows(df, index=False, header=True):
            ws.append(row)

        for column_cells in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
            col_letter = get_column_letter(column_cells[0].column)
            ws.column_dimensions[col_letter].width = max_length + 2

        wb.save(excel_file)
        print(f"Excel file created: {excel_file}")

    except FileNotFoundError:
        print(f"The file {json_file} was not found.")
    except json.JSONDecodeError:
        print(f"Error decoding JSON from the file {json_file}.")
    except Exception as e:
        print(f"An error occurred: {e}")
